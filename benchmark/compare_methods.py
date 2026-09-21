"""Create a fair comparison workbook for proposed DDQN versus benchmark.

Only physically commensurable metrics are compared. Raw reward, Q-values,
benchmark-specific discomfort, and raw household utility are intentionally
excluded.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


PERIOD_ORDER = ["july", "august", "november"]


def resolve_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        sampled_max_row = min(sheet.max_row, 250)
        for column in sheet.iter_cols(
            min_row=1, max_row=sampled_max_row, min_col=1, max_col=sheet.max_column
        ):
            width = min(
                max(len(str(cell.value)) if cell.value is not None else 0 for cell in column) + 2,
                48,
            )
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(path)


def prepare_proposed(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path, sheet_name="Seed_Performance")
    required = {"period", "seed"}
    if not required.issubset(frame.columns):
        raise ValueError(f"Proposed summary is missing {sorted(required - set(frame.columns))}")
    result = pd.DataFrame({"period": frame["period"].str.lower(), "seed": frame["seed"], "method": "Proposed"})
    direct = {
        "peak_reduction_percent": "peak_reduction_percent",
        "PAR_reduction_percent": "PAR_reduction_percent",
        "remaining_violation_rate_percent": "remaining_violation_rate_percent",
        "cumulative_excess_kWh": "cumulative_excess_kWh",
        "maximum_overrun_kW": "maximum_overrun_kW",
        "tracking_MAE_kWh": "external_tracking_MAE_kWh",
        "useful_DR_kWh": "useful_DR_kWh",
        "excess_DR_kWh": "excess_DR_kWh",
        "incentive_payment_cent": "incentive_payment_cent",
        "unnecessary_incentive_rate_percent": "unnecessary_intervention_rate_percent_external",
        "no_response_rate_percent": "no_response_rate_percent",
    }
    for source, target in direct.items():
        if source in frame:
            result[target] = pd.to_numeric(frame[source], errors="coerce")
    if {"net_wholesale_procurement_impact_cent", "incentive_payment_cent"}.issubset(frame.columns):
        result["net_wholesale_after_payment_cent"] = (
            pd.to_numeric(frame["net_wholesale_procurement_impact_cent"], errors="coerce")
            - pd.to_numeric(frame["incentive_payment_cent"], errors="coerce")
        )
    total_response = (
        pd.to_numeric(frame.get("useful_DR_kWh"), errors="coerce")
        + pd.to_numeric(frame.get("excess_DR_kWh"), errors="coerce")
    )
    result["useful_DR_ratio_percent"] = 100.0 * pd.to_numeric(
        frame.get("useful_DR_kWh"), errors="coerce"
    ) / total_response.replace(0, np.nan)
    result["payment_per_useful_kWh_cent"] = pd.to_numeric(
        frame.get("incentive_payment_cent"), errors="coerce"
    ) / pd.to_numeric(frame.get("useful_DR_kWh"), errors="coerce").replace(0, np.nan)
    if "rebound_violation_hours" in frame:
        result["rebound_violation_hours"] = pd.to_numeric(
            frame["rebound_violation_hours"], errors="coerce"
        )
    return result


def prepare_benchmark(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path, sheet_name="Seed_Performance")
    result = frame.copy()
    result["period"] = result["period"].str.lower()
    result["method"] = "Benchmark"
    keep = [
        "period",
        "seed",
        "method",
        "peak_reduction_percent",
        "PAR_reduction_percent",
        "mean_load_reduction_percent",
        "remaining_violation_rate_percent",
        "capacity_compliance_improvement_percentage_points",
        "cumulative_excess_kWh",
        "maximum_overrun_kW",
        "external_tracking_MAE_kWh",
        "total_curtailment_kWh",
        "useful_DR_kWh",
        "excess_DR_kWh",
        "useful_DR_ratio_percent",
        "overload_targeting_ratio_percent",
        "incentive_payment_cent",
        "payment_per_useful_kWh_cent",
        "net_wholesale_after_payment_cent",
        "unnecessary_intervention_rate_percent_external",
        "no_response_rate_percent",
    ]
    return result[[column for column in keep if column in result.columns]]


def long_mean_sd(combined: pd.DataFrame) -> pd.DataFrame:
    id_columns = {"period", "seed", "method"}
    numeric = [column for column in combined.select_dtypes(include=[np.number]).columns if column not in id_columns]
    rows = []
    for (period, method), frame in combined.groupby(["period", "method"], sort=False, observed=True):
        for metric in numeric:
            values = pd.to_numeric(frame[metric], errors="coerce").dropna()
            if values.empty:
                continue
            rows.append(
                {
                    "period": period,
                    "method": method,
                    "metric": metric,
                    "mean": float(values.mean()),
                    "sample_SD": float(values.std(ddof=1)) if len(values) > 1 else np.nan,
                    "n_seeds": int(len(values)),
                }
            )
    return pd.DataFrame(rows)


def differences(mean_sd: pd.DataFrame) -> pd.DataFrame:
    pivot = mean_sd.pivot_table(index=["period", "metric"], columns="method", values="mean").reset_index()
    if {"Proposed", "Benchmark"}.issubset(pivot.columns):
        pivot["proposed_minus_benchmark"] = pivot["Proposed"] - pivot["Benchmark"]
    return pivot


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--proposed_summary", default="results/summary.xlsx")
    parser.add_argument("--benchmark_summary", default="results/benchmark/summary.xlsx")
    parser.add_argument("--output", default="results/method_comparison.xlsx")
    args = parser.parse_args()

    proposed = prepare_proposed(resolve_path(args.proposed_summary))
    benchmark = prepare_benchmark(resolve_path(args.benchmark_summary))
    combined = pd.concat([proposed, benchmark], ignore_index=True, sort=False)
    combined["period"] = pd.Categorical(combined["period"], PERIOD_ORDER, ordered=True)
    combined = combined.sort_values(["period", "method", "seed"]).reset_index(drop=True)
    mean_sd = long_mean_sd(combined)
    diff = differences(mean_sd)

    structural = pd.DataFrame(
        [
            ("Capacity observed by policy", "No", "Yes"),
            ("Capacity used in reward", "No", "Yes"),
            ("Appliance-level feasibility", "No", "Yes"),
            ("Time shifting", "No", "Yes"),
            ("Response type", "Curtailment only", "Curtailment plus feasible shifting"),
            ("Rebound metric", "N/A - structurally undefined", "Measured"),
            ("External 7 kW evaluation", "Yes, post hoc only", "Yes"),
        ],
        columns=["feature", "Benchmark", "Proposed"],
    )
    exclusions = pd.DataFrame(
        [
            ("Raw reward", "Different objective functions and scales"),
            ("Q-values", "Different state and reward definitions"),
            ("Raw discomfort", "Different discomfort models"),
            ("Raw household utility", "Different mathematical definitions"),
        ],
        columns=["excluded_quantity", "reason"],
    )

    output = resolve_path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        combined.to_excel(writer, sheet_name="Seed_Comparison", index=False)
        mean_sd.to_excel(writer, sheet_name="Mean_SD_Comparison", index=False)
        diff.to_excel(writer, sheet_name="Mean_Differences", index=False)
        structural.to_excel(writer, sheet_name="Structural_Comparison", index=False)
        exclusions.to_excel(writer, sheet_name="Excluded_Quantities", index=False)
    style_workbook(output)
    print(f"Method comparison saved: {output}")


if __name__ == "__main__":
    main()
