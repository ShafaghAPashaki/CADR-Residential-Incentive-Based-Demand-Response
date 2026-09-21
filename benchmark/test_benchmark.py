"""Test one trained benchmark checkpoint over a complete evaluation period.

Capacity is used only after each capacity-blind trajectory is complete, as an
external operational reference for descriptive metrics and plots.
"""

from __future__ import annotations

import argparse
import json
import math
import shutil
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from agent.agent_ddqn import DDQNAgent
from benchmark.env_benchmark import Environment
from utils.config_loader import load_config


SCRIPT_VERSION = "benchmark_test_v2"
EPS = 1.0e-9


def safe_div(numerator: float, denominator: float) -> float:
    return float(numerator / denominator) if abs(float(denominator)) > EPS else float("nan")


def percent_reduction(before: float, after: float) -> float:
    return 100.0 * safe_div(before - after, before)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sampled_max_row = min(sheet.max_row, 250)
        for column in sheet.iter_cols(
            min_row=1, max_row=sampled_max_row, min_col=1, max_col=sheet.max_column
        ):
            letter = column[0].column_letter
            width = min(
                max(len(str(cell.value)) if cell.value is not None else 0 for cell in column) + 2,
                45,
            )
            sheet.column_dimensions[letter].width = max(width, 12)
    workbook.save(path)


class BenchmarkTester:
    def __init__(
        self,
        model_path: Path,
        cfg: dict[str, Any],
        period: str,
        start_day: int,
        end_day: int,
        plot_days: list[int],
        output_dir: Path,
        overwrite: bool = False,
        make_plots: bool = True,
    ) -> None:
        self.model_path = model_path.resolve()
        self.cfg = cfg
        self.period = str(period).lower()
        self.start_day = int(start_day)
        self.end_day = int(end_day)
        self.plot_days = [int(value) for value in plot_days]
        self.output_dir = output_dir.resolve()
        self.make_plots = bool(make_plots)
        self.seed = int(cfg["general"]["seed"])
        self.capacity = float(cfg["external_reference"]["capacity_threshold_kw"])
        self.house_ids = [int(value) for value in cfg["environment"]["house_ids"]]

        if self.end_day < self.start_day:
            raise ValueError("end_day must be greater than or equal to start_day.")
        invalid_plots = [day for day in self.plot_days if not self.start_day <= day <= self.end_day]
        if invalid_plots:
            raise ValueError(f"Plot days fall outside the test period: {invalid_plots}")

        if self.output_dir.exists() and any(self.output_dir.iterdir()):
            if not overwrite:
                raise FileExistsError(
                    f"Output directory is non-empty: {self.output_dir}. "
                    "Use --overwrite only when intentionally replacing this test output."
                )
            shutil.rmtree(self.output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.env = Environment(self.house_ids, cfg_override=cfg, rng_stream=200_000)
        state = self.env.reset(day=self.start_day, mode="test")
        self.agent = DDQNAgent(
            len(state),
            self.env.num_actions,
            cfg_override=cfg,
            action_costs=np.mean(self.env.all_actions, axis=1),
        )
        checkpoint = self.agent.load(str(self.model_path), load_optimizer=False)
        self._validate_checkpoint(checkpoint)
        self.agent.policy_net.eval()
        self.checkpoint = checkpoint

    def _validate_checkpoint(self, checkpoint: dict[str, Any]) -> None:
        required = {
            "environment_version": self.env.VERSION,
            "accounting_version": self.env.ACCOUNTING_VERSION,
            "response_model_version": self.env.RESPONSE_MODEL_VERSION,
            "reward_version": self.env.REWARD_VERSION,
            "config_signature": self.env.get_config_signature(),
            "seed": self.seed,
            "capacity_blind": True,
            "curtailment_only": True,
        }
        mismatches = []
        for key, expected in required.items():
            observed = checkpoint.get(key)
            if observed != expected:
                mismatches.append(f"{key}: checkpoint={observed!r}, expected={expected!r}")
        if int(checkpoint.get("state_dim", -1)) != self.env.state_dim:
            mismatches.append(
                f"state_dim: checkpoint={checkpoint.get('state_dim')}, expected={self.env.state_dim}"
            )
        if int(checkpoint.get("action_dim", -1)) != self.env.num_actions:
            mismatches.append(
                f"action_dim: checkpoint={checkpoint.get('action_dim')}, expected={self.env.num_actions}"
            )
        if checkpoint.get("state_feature_names") != self.env.get_state_feature_names():
            mismatches.append("state_feature_names do not match the active benchmark environment")
        if [int(v) for v in checkpoint.get("house_ids", [])] != self.house_ids:
            mismatches.append("house_ids do not match the active benchmark config")
        if mismatches:
            raise ValueError("Checkpoint provenance validation failed:\n- " + "\n- ".join(mismatches))

    def run(self) -> None:
        aggregate_rows: list[dict[str, Any]] = []
        household_rows: list[dict[str, Any]] = []
        plot_cache: dict[int, dict[str, np.ndarray]] = {}

        for day in range(self.start_day, self.end_day + 1):
            state = self.env.reset(day=day, mode="test")
            done = False
            step = 0
            while not done:
                action_index = self.agent.greedy_action(state)
                next_state, reward, done, _ = self.env.step(action_index)
                h = self.env.curr_step - 1
                timestamp = self.env.price_df.index[h]

                baseline = self.env.baseline_per_house[h].copy()
                reduction = self.env.reductions[h].copy()
                post = self.env.after_per_house[h].copy()
                raw = self.env.raw_actions[h].copy()
                nominal = self.env.incentives[h].copy()
                incremental = self.env.incremental_incentives[h].copy()
                active_rate = np.where(raw > EPS, nominal, 0.0)
                payment = self.env.incentive_payment_per_house[h].copy()
                discomfort = self.env.discomforts[h].copy()
                eu_reward = self.env.rewards_customers[h].copy()

                baseline_total = float(baseline.sum())
                reduction_total = float(reduction.sum())
                post_total = float(post.sum())
                required = max(baseline_total - self.capacity, 0.0)
                useful = min(reduction_total, required)
                excess = max(reduction_total - required, 0.0)
                unmet = max(required - reduction_total, 0.0)
                overrun = max(post_total - self.capacity, 0.0)
                intervention = bool(np.any(raw > EPS))
                no_need = bool(baseline_total <= self.capacity + EPS)

                aggregate = {
                    "period": self.period,
                    "seed": self.seed,
                    "day": int(day),
                    "timestamp": timestamp,
                    "hour": int(self.env.hours[h]),
                    "price_cent_per_kWh": float(self.env.prices[h]),
                    "elasticity": float(self.env.elasticities[h]),
                    "daily_reference_price_cent_per_kWh": float(self.env.daily_reference_price),
                    "lambda_min_cent_per_kWh": float(self.env.lambda_min[h]),
                    "lambda_max_cent_per_kWh": float(self.env.lambda_max[h]),
                    "action_index": int(action_index),
                    "baseline_load_kW": baseline_total,
                    "post_DR_load_kW": post_total,
                    "total_curtailment_kWh": reduction_total,
                    "intervention_flag": int(intervention),
                    "no_need_flag_external": int(no_need),
                    "baseline_violation_flag_external": int(baseline_total > self.capacity + EPS),
                    "post_violation_flag_external": int(overrun > EPS),
                    "required_reduction_kWh_external": required,
                    "useful_DR_kWh_external": useful,
                    "excess_DR_kWh_external": excess,
                    "unmet_reduction_kWh_external": unmet,
                    "post_overrun_kW_external": overrun,
                    "unnecessary_reduction_kWh_external": reduction_total if no_need else 0.0,
                    "incentive_payment_cent": float(payment.sum()),
                    "avoided_wholesale_value_cent": float(self.env.wholesale_avoided_value[h]),
                    "net_wholesale_after_payment_cent": float(self.env.rewards_service_provider[h]),
                    "benchmark_customer_reward_cent": float(eu_reward.sum()),
                    "benchmark_total_reward": float(reward),
                }
                for idx, house_id in enumerate(self.house_ids):
                    aggregate[f"raw_action_house_{house_id}"] = float(raw[idx])
                    aggregate[f"active_nominal_rate_house_{house_id}_cent_per_kWh"] = float(active_rate[idx])
                    aggregate[f"baseline_house_{house_id}_kW"] = float(baseline[idx])
                    aggregate[f"curtailment_house_{house_id}_kWh"] = float(reduction[idx])
                    aggregate[f"post_house_{house_id}_kW"] = float(post[idx])
                aggregate_rows.append(aggregate)

                for idx, house_id in enumerate(self.house_ids):
                    expected_reduction = min(
                        baseline[idx] * self.env.elasticities[h] * raw[idx],
                        self.env.max_reduction_fraction * baseline[idx],
                    ) if self.env.daily_reference_price > 0.0 else 0.0
                    expected_nominal = self.env.lambda_min[h] + raw[idx] * (
                        self.env.lambda_max[h] - self.env.lambda_min[h]
                    )
                    expected_payment = nominal[idx] * reduction[idx]
                    expected_discomfort = (
                        0.5 * self.env.mu[idx] * reduction[idx] ** 2
                        + self.env.kappa * reduction[idx]
                    )
                    expected_eu = (
                        self.env.rho * expected_payment
                        - (1.0 - self.env.rho) * expected_discomfort
                    )
                    household_rows.append(
                        {
                            "period": self.period,
                            "seed": self.seed,
                            "day": int(day),
                            "timestamp": timestamp,
                            "hour": int(self.env.hours[h]),
                            "house_id": int(house_id),
                            "action_index": int(action_index),
                            "raw_action": float(raw[idx]),
                            "intervention_flag": int(raw[idx] > EPS),
                            "nominal_incentive_rate_cent_per_kWh": float(nominal[idx]),
                            "active_nominal_rate_cent_per_kWh": float(active_rate[idx]),
                            "incremental_incentive_rate_cent_per_kWh": float(incremental[idx]),
                            "baseline_kW": float(baseline[idx]),
                            "curtailment_kWh": float(reduction[idx]),
                            "post_DR_kW": float(post[idx]),
                            "payment_cent": float(payment[idx]),
                            "benchmark_discomfort_cent": float(discomfort[idx]),
                            "benchmark_EU_reward_cent": float(eu_reward[idx]),
                            "no_response_flag": int(raw[idx] > EPS and reduction[idx] <= EPS),
                            "response_reconstruction_error": float(reduction[idx] - expected_reduction),
                            "nominal_rate_reconstruction_error": float(nominal[idx] - expected_nominal),
                            "payment_reconstruction_error": float(payment[idx] - expected_payment),
                            "discomfort_reconstruction_error": float(discomfort[idx] - expected_discomfort),
                            "EU_reward_reconstruction_error": float(eu_reward[idx] - expected_eu),
                        }
                    )

                state = next_state
                step += 1
                if step > self.env.max_steps:
                    raise RuntimeError(f"Day {day} exceeded the configured episode length.")

            if day in self.plot_days:
                plot_cache[day] = {
                    "baseline": self.env.baseline_per_house.copy(),
                    "reduction": self.env.reductions.copy(),
                    "post": self.env.after_per_house.copy(),
                    "raw": self.env.raw_actions.copy(),
                    "active_rate": np.where(self.env.raw_actions > EPS, self.env.incentives, 0.0),
                    "price": self.env.prices.copy(),
                }

        aggregate_df = pd.DataFrame(aggregate_rows)
        household_df = pd.DataFrame(household_rows)
        daily_df = self._daily_metrics(aggregate_df, household_df)
        monthly_df = self._monthly_metrics(aggregate_df, household_df, daily_df)
        household_summary_df = self._household_summary(household_df)
        economics_df = self._economics_summary(aggregate_df, household_df)
        external_df = self._external_capacity_summary(monthly_df)
        audit_summary_df, aggregate_audit_df = self._audit_summary(aggregate_df, household_df)
        metadata_df = self._metadata_frame()
        definitions_df = self._metric_definitions()

        results_path = self.output_dir / "results.xlsx"
        with pd.ExcelWriter(results_path, engine="openpyxl") as writer:
            aggregate_df.to_excel(writer, sheet_name="Hourly_Aggregate", index=False)
            household_df.to_excel(writer, sheet_name="Hourly_Household", index=False)
            daily_df.to_excel(writer, sheet_name="Daily", index=False)
            monthly_df.to_excel(writer, sheet_name="Monthly", index=False)
            household_summary_df.to_excel(writer, sheet_name="Household", index=False)
            economics_df.to_excel(writer, sheet_name="Economics", index=False)
            external_df.to_excel(writer, sheet_name="Capacity_External", index=False)
            definitions_df.to_excel(writer, sheet_name="Metric_Definitions", index=False)
            metadata_df.to_excel(writer, sheet_name="Run_Metadata", index=False)
        style_workbook(results_path)

        audit_path = self.output_dir / "audit.xlsx"
        with pd.ExcelWriter(audit_path, engine="openpyxl") as writer:
            household_df.to_excel(writer, sheet_name="Household_Ledger", index=False)
            aggregate_audit_df.to_excel(writer, sheet_name="Aggregate_Reconstruction", index=False)
            audit_summary_df.to_excel(writer, sheet_name="Audit_Summary", index=False)
            metadata_df.to_excel(writer, sheet_name="Run_Metadata", index=False)
        style_workbook(audit_path)

        if self.make_plots:
            for day, arrays in plot_cache.items():
                self._plot_aggregated(day, arrays)
                self._plot_reduction_alignment(day, arrays)
                self._plot_households(day, arrays)

        print(
            f"Benchmark test complete | period={self.period} | seed={self.seed} | "
            f"audit={audit_path} | results={results_path}"
        )

    def _daily_metrics(self, agg: pd.DataFrame, hh: pd.DataFrame) -> pd.DataFrame:
        rows = []
        for day, frame in agg.groupby("day", sort=True):
            hh_day = hh[hh["day"] == day]
            rows.append(self._summarize_period(frame, hh_day, day=int(day)))
        return pd.DataFrame(rows)

    def _monthly_metrics(
        self, agg: pd.DataFrame, hh: pd.DataFrame, daily: pd.DataFrame
    ) -> pd.DataFrame:
        row = self._summarize_period(agg, hh, day=None)
        # The proposed-method summary defines monthly peak/PAR performance from
        # the mean of daily peak and daily PAR values. Preserve global extrema
        # as supplementary fields, but use the identical daily-mean definition
        # for the primary cross-method columns.
        row["global_month_peak_baseline_kW"] = row["peak_baseline_kW"]
        row["global_month_peak_post_DR_kW"] = row["peak_post_DR_kW"]
        row["global_month_peak_reduction_percent"] = row["peak_reduction_percent"]
        row["peak_baseline_kW"] = float(daily["peak_baseline_kW"].mean())
        row["peak_post_DR_kW"] = float(daily["peak_post_DR_kW"].mean())
        row["peak_reduction_percent"] = percent_reduction(
            row["peak_baseline_kW"], row["peak_post_DR_kW"]
        )
        row["PAR_baseline"] = float(daily["PAR_baseline"].mean())
        row["PAR_post_DR"] = float(daily["PAR_post_DR"].mean())
        row["PAR_reduction_percent"] = percent_reduction(
            row["PAR_baseline"], row["PAR_post_DR"]
        )
        return pd.DataFrame([row])

    def _summarize_period(
        self, agg: pd.DataFrame, hh: pd.DataFrame, day: int | None
    ) -> dict[str, Any]:
        baseline = agg["baseline_load_kW"].to_numpy(dtype=float)
        post = agg["post_DR_load_kW"].to_numpy(dtype=float)
        peak_baseline = float(np.max(baseline))
        peak_post = float(np.max(post))
        mean_baseline = float(np.mean(baseline))
        mean_post = float(np.mean(post))
        par_baseline = safe_div(peak_baseline, mean_baseline)
        par_post = safe_div(peak_post, mean_post)
        total_reduction = float(agg["total_curtailment_kWh"].sum())
        intervention_hours = int(agg["intervention_flag"].sum())
        hh_interventions = int(hh["intervention_flag"].sum())
        household_no_response = int(hh["no_response_flag"].sum())
        no_response_hours = int(
            ((agg["intervention_flag"] > 0) & (agg["total_curtailment_kWh"] <= EPS)).sum()
        )
        no_need_hours = int(agg["no_need_flag_external"].sum())
        unnecessary_hours = int(
            ((agg["intervention_flag"] > 0) & (agg["no_need_flag_external"] > 0)).sum()
        )
        useful = float(agg["useful_DR_kWh_external"].sum())
        excess = float(agg["excess_DR_kWh_external"].sum())
        overload_reduction = float(
            agg.loc[agg["baseline_violation_flag_external"] > 0, "total_curtailment_kWh"].sum()
        )
        baseline_violation_hours = int(agg["baseline_violation_flag_external"].sum())
        post_violation_hours = int(agg["post_violation_flag_external"].sum())
        hours = int(len(agg))
        payment = float(agg["incentive_payment_cent"].sum())
        avoided = float(agg["avoided_wholesale_value_cent"].sum())
        result = {
            "period": self.period,
            "seed": self.seed,
            "day": day if day is not None else "all",
            "days": int(agg["day"].nunique()),
            "hours": hours,
            "peak_baseline_kW": peak_baseline,
            "peak_post_DR_kW": peak_post,
            "peak_reduction_percent": percent_reduction(peak_baseline, peak_post),
            "mean_baseline_kW": mean_baseline,
            "mean_post_DR_kW": mean_post,
            "mean_load_reduction_percent": percent_reduction(mean_baseline, mean_post),
            "PAR_baseline": par_baseline,
            "PAR_post_DR": par_post,
            "PAR_reduction_percent": percent_reduction(par_baseline, par_post),
            "total_curtailment_kWh": total_reduction,
            "intervention_hours": intervention_hours,
            "household_intervention_events": hh_interventions,
            "no_response_intervention_hours": no_response_hours,
            "no_response_rate_percent": 100.0 * safe_div(no_response_hours, intervention_hours),
            "household_no_response_events": household_no_response,
            "household_no_response_rate_percent": 100.0 * safe_div(
                household_no_response, hh_interventions
            ),
            "no_need_hours_external": no_need_hours,
            "unnecessary_intervention_hours_external": unnecessary_hours,
            "unnecessary_intervention_rate_percent_external": 100.0 * safe_div(
                unnecessary_hours, no_need_hours
            ),
            "unnecessary_reduction_kWh_external": float(
                agg["unnecessary_reduction_kWh_external"].sum()
            ),
            "incentive_payment_cent": payment,
            "avoided_wholesale_value_cent": avoided,
            "net_wholesale_after_payment_cent": avoided - payment,
            "benchmark_SP_reward_cent": float(agg["net_wholesale_after_payment_cent"].sum()),
            "benchmark_customer_reward_cent": float(agg["benchmark_customer_reward_cent"].sum()),
            "benchmark_total_reward": float(agg["benchmark_total_reward"].sum()),
            "baseline_violation_hours_external": baseline_violation_hours,
            "post_violation_hours_external": post_violation_hours,
            "baseline_violation_rate_percent_external": 100.0 * safe_div(
                baseline_violation_hours, hours
            ),
            "remaining_violation_rate_percent": 100.0 * safe_div(post_violation_hours, hours),
            "capacity_compliance_improvement_percentage_points": 100.0
            * safe_div(baseline_violation_hours - post_violation_hours, hours),
            "cumulative_excess_kWh": float(agg["post_overrun_kW_external"].sum()),
            "maximum_overrun_kW": float(agg["post_overrun_kW_external"].max()),
            "external_tracking_MAE_kWh": float(
                np.mean(
                    np.abs(
                        agg["total_curtailment_kWh"].to_numpy(dtype=float)
                        - agg["required_reduction_kWh_external"].to_numpy(dtype=float)
                    )
                )
            ),
            "required_reduction_kWh_external": float(
                agg["required_reduction_kWh_external"].sum()
            ),
            "useful_DR_kWh": useful,
            "excess_DR_kWh": excess,
            "unmet_reduction_kWh_external": float(
                agg["unmet_reduction_kWh_external"].sum()
            ),
            "useful_DR_ratio_percent": 100.0 * safe_div(useful, total_reduction),
            "overload_targeting_ratio_percent": 100.0
            * safe_div(overload_reduction, total_reduction),
            "payment_per_useful_kWh_cent": safe_div(payment, useful),
            "rebound_status": "N/A - curtailment-only model",
        }
        return result

    def _household_summary(self, hh: pd.DataFrame) -> pd.DataFrame:
        rows = []
        for house_id, frame in hh.groupby("house_id", sort=False):
            interventions = int(frame["intervention_flag"].sum())
            no_response = int(frame["no_response_flag"].sum())
            rows.append(
                {
                    "period": self.period,
                    "seed": self.seed,
                    "house_id": int(house_id),
                    "baseline_energy_kWh": float(frame["baseline_kW"].sum()),
                    "post_DR_energy_kWh": float(frame["post_DR_kW"].sum()),
                    "total_curtailment_kWh": float(frame["curtailment_kWh"].sum()),
                    "intervention_events": interventions,
                    "no_response_events": no_response,
                    "no_response_rate_percent": 100.0 * safe_div(no_response, interventions),
                    "incentive_payment_cent": float(frame["payment_cent"].sum()),
                    "benchmark_discomfort_cent": float(frame["benchmark_discomfort_cent"].sum()),
                    "benchmark_EU_reward_cent": float(frame["benchmark_EU_reward_cent"].sum()),
                }
            )
        return pd.DataFrame(rows)

    def _economics_summary(self, agg: pd.DataFrame, hh: pd.DataFrame) -> pd.DataFrame:
        rows = [
            {
                "entity": "Service provider",
                "metric": "Avoided wholesale procurement value",
                "value_cent": float(agg["avoided_wholesale_value_cent"].sum()),
                "cross_method_comparable": True,
            },
            {
                "entity": "Service provider",
                "metric": "Incentive payment",
                "value_cent": float(agg["incentive_payment_cent"].sum()),
                "cross_method_comparable": True,
            },
            {
                "entity": "Service provider",
                "metric": "Net wholesale impact after payment",
                "value_cent": float(agg["net_wholesale_after_payment_cent"].sum()),
                "cross_method_comparable": True,
            },
        ]
        for house_id, frame in hh.groupby("house_id", sort=False):
            rows.extend(
                [
                    {
                        "entity": f"Household {int(house_id)}",
                        "metric": "Incentive income",
                        "value_cent": float(frame["payment_cent"].sum()),
                        "cross_method_comparable": True,
                    },
                    {
                        "entity": f"Household {int(house_id)}",
                        "metric": "Benchmark-specific discomfort",
                        "value_cent": float(frame["benchmark_discomfort_cent"].sum()),
                        "cross_method_comparable": False,
                    },
                    {
                        "entity": f"Household {int(house_id)}",
                        "metric": "Benchmark-specific EU reward",
                        "value_cent": float(frame["benchmark_EU_reward_cent"].sum()),
                        "cross_method_comparable": False,
                    },
                ]
            )
        return pd.DataFrame(rows)

    def _external_capacity_summary(self, monthly: pd.DataFrame) -> pd.DataFrame:
        columns = [
            "period",
            "seed",
            "baseline_violation_rate_percent_external",
            "remaining_violation_rate_percent",
            "capacity_compliance_improvement_percentage_points",
            "cumulative_excess_kWh",
            "maximum_overrun_kW",
            "external_tracking_MAE_kWh",
            "required_reduction_kWh_external",
            "useful_DR_kWh",
            "excess_DR_kWh",
            "unmet_reduction_kWh_external",
            "useful_DR_ratio_percent",
            "overload_targeting_ratio_percent",
            "payment_per_useful_kWh_cent",
            "unnecessary_intervention_hours_external",
            "unnecessary_reduction_kWh_external",
        ]
        frame = monthly[columns].copy()
        frame.insert(2, "capacity_threshold_kW", self.capacity)
        frame["capacity_available_to_policy"] = False
        frame["capacity_usage"] = "post-hoc evaluation only"
        return frame

    def _audit_summary(
        self, agg: pd.DataFrame, hh: pd.DataFrame
    ) -> tuple[pd.DataFrame, pd.DataFrame]:
        aggregate_audit = agg[
            [
                "period",
                "seed",
                "day",
                "timestamp",
                "hour",
                "benchmark_SP_reward_cent",
                "benchmark_customer_reward_cent",
                "benchmark_total_reward",
            ]
        ].copy() if "benchmark_SP_reward_cent" in agg.columns else agg[["period", "seed", "day", "timestamp", "hour"]].copy()

        expected_sp = agg["avoided_wholesale_value_cent"] - agg["incentive_payment_cent"]
        expected_total = (
            self.env.alpha * expected_sp
            + (1.0 - self.env.alpha) * agg["benchmark_customer_reward_cent"]
        )
        aggregate_audit["SP_reward_reconstruction_error"] = (
            agg["net_wholesale_after_payment_cent"] - expected_sp
        )
        aggregate_audit["total_reward_reconstruction_error"] = (
            agg["benchmark_total_reward"] - expected_total
        )
        aggregate_audit["energy_balance_error"] = (
            agg["baseline_load_kW"]
            - agg["post_DR_load_kW"]
            - agg["total_curtailment_kWh"]
        )

        metrics = {
            "max_abs_response_reconstruction_error": float(
                hh["response_reconstruction_error"].abs().max()
            ),
            "max_abs_nominal_rate_reconstruction_error": float(
                hh["nominal_rate_reconstruction_error"].abs().max()
            ),
            "max_abs_payment_reconstruction_error": float(
                hh["payment_reconstruction_error"].abs().max()
            ),
            "max_abs_discomfort_reconstruction_error": float(
                hh["discomfort_reconstruction_error"].abs().max()
            ),
            "max_abs_EU_reward_reconstruction_error": float(
                hh["EU_reward_reconstruction_error"].abs().max()
            ),
            "max_abs_SP_reward_reconstruction_error": float(
                aggregate_audit["SP_reward_reconstruction_error"].abs().max()
            ),
            "max_abs_total_reward_reconstruction_error": float(
                aggregate_audit["total_reward_reconstruction_error"].abs().max()
            ),
            "max_abs_energy_balance_error": float(
                aggregate_audit["energy_balance_error"].abs().max()
            ),
            "invalid_negative_reduction_count": int((hh["curtailment_kWh"] < -EPS).sum()),
            "invalid_post_load_count": int((hh["post_DR_kW"] < -EPS).sum()),
            "zero_action_nonzero_response_count": int(
                ((hh["raw_action"].abs() <= EPS) & (hh["curtailment_kWh"].abs() > EPS)).sum()
            ),
            "capacity_used_by_policy": False,
        }
        audit_summary = pd.DataFrame(
            [{"metric": key, "value": value} for key, value in metrics.items()]
        )
        return audit_summary, aggregate_audit

    def _metadata_frame(self) -> pd.DataFrame:
        values = {
            "script_version": SCRIPT_VERSION,
            "model_path": str(self.model_path),
            "period": self.period,
            "seed": self.seed,
            "test_range": f"{self.start_day}-{self.end_day}",
            "plot_days": json.dumps(self.plot_days),
            "capacity_threshold_kW_external": self.capacity,
            "capacity_available_to_policy": False,
            "environment_version": self.env.VERSION,
            "accounting_version": self.env.ACCOUNTING_VERSION,
            "response_model_version": self.env.RESPONSE_MODEL_VERSION,
            "reward_version": self.env.REWARD_VERSION,
            "negative_price_convention": self.env.NEGATIVE_PRICE_CONVENTION,
            "config_signature": self.env.get_config_signature(),
            "checkpoint_episode": self.checkpoint.get("episode"),
            "checkpoint_model_type": self.checkpoint.get("model_type"),
            "best_validation_episode": self.checkpoint.get("best_validation_episode"),
            "state_dim": self.env.state_dim,
            "state_features": json.dumps(self.env.get_state_feature_names()),
            "action_dim": self.env.num_actions,
            "house_ids": json.dumps(self.house_ids),
            "rebound": "N/A - curtailment-only response model",
        }
        return pd.DataFrame([{"field": key, "value": value} for key, value in values.items()])

    @staticmethod
    def _metric_definitions() -> pd.DataFrame:
        rows = [
            ("Peak reduction", "100*(baseline peak - post-DR peak)/baseline peak", "native", True),
            ("PAR reduction", "100*(baseline PAR - post-DR PAR)/baseline PAR", "native", True),
            ("Total curtailment", "Sum of baseline minus post-DR load", "native", True),
            ("Intervention", "At least one household raw action is greater than zero", "native", True),
            ("No-response rate", "Aggregate intervention hours with zero aggregate curtailment divided by aggregate intervention hours", "native", True),
            ("Household no-response event", "Household raw action > 0 and realised household curtailment = 0", "native", False),
            ("Unnecessary-intervention rate", "Intervention hours in external no-need states divided by all external no-need hours", "external", True),
            ("Net wholesale after payment", "price*curtailment - incentive payment", "native", True),
            ("Required reduction", "max(baseline aggregate load - 7 kW, 0)", "external", True),
            ("Useful DR", "min(total curtailment, required reduction)", "external", True),
            ("Excess DR", "max(total curtailment - required reduction, 0)", "external", True),
            ("Useful-DR ratio", "100*useful DR/total curtailment", "external", True),
            ("Overload-targeting ratio", "100*curtailment in baseline-overload hours/total curtailment", "external", True),
            ("Payment per useful kWh", "total incentive payment/useful DR", "external", True),
            ("Benchmark discomfort", "0.5*mu*curtailment^2 + kappa*curtailment", "native", False),
            ("Benchmark total reward", "alpha*R_SP + (1-alpha)*sum(R_EU)", "native", False),
            ("Rebound", "Not applicable because benchmark curtails and never shifts energy", "structural", False),
        ]
        return pd.DataFrame(
            rows,
            columns=["metric", "definition", "category", "cross_method_comparable"],
        )

    def _plot_aggregated(self, day: int, arrays: dict[str, np.ndarray]) -> None:
        baseline = arrays["baseline"].sum(axis=1)
        post = arrays["post"].sum(axis=1)
        hours = np.arange(24)
        plt.figure(figsize=(11, 6))
        plt.plot(hours, baseline, marker="o", linewidth=2, label="Baseline load")
        plt.plot(hours, post, marker="o", linestyle="--", linewidth=2, label="Benchmark post-DR load")
        plt.axhline(
            self.capacity,
            linestyle=":",
            linewidth=2,
            label=f"External capacity reference ({self.capacity:g} kW)",
        )
        plt.xlabel("Hour")
        plt.ylabel("Aggregated load (kW)")
        plt.xticks(np.arange(0, 24, 2))
        plt.xlim(0, 23)
        plt.legend(fontsize=10)
        plt.tight_layout()
        plt.savefig(self.output_dir / f"aggregated_load_day_{day}.png", dpi=300, bbox_inches="tight")
        plt.close()

    def _plot_reduction_alignment(self, day: int, arrays: dict[str, np.ndarray]) -> None:
        baseline = arrays["baseline"].sum(axis=1)
        reduction = arrays["reduction"].sum(axis=1)
        required = np.maximum(baseline - self.capacity, 0.0)
        hours = np.arange(24)
        plt.figure(figsize=(11, 6))
        plt.step(hours, required, where="mid", linewidth=2, label="Required reduction (external)")
        plt.step(hours, reduction, where="mid", linewidth=2, label="Benchmark curtailment")
        plt.xlabel("Hour")
        plt.ylabel("Reduction (kWh)")
        plt.xticks(np.arange(0, 24, 2))
        plt.xlim(0, 23)
        plt.legend(fontsize=10)
        plt.tight_layout()
        plt.savefig(
            self.output_dir / f"reduction_alignment_day_{day}.png",
            dpi=300,
            bbox_inches="tight",
        )
        plt.close()

    def _plot_households(self, day: int, arrays: dict[str, np.ndarray]) -> None:
        hours = np.arange(24)
        for idx, house_id in enumerate(self.house_ids):
            fig, left = plt.subplots(figsize=(11, 6))
            left.plot(hours, arrays["baseline"][:, idx], linewidth=2, label="Baseline")
            left.plot(
                hours, arrays["post"][:, idx], linestyle="--", linewidth=2, label="Post-DR"
            )
            left.set_xlabel("Hour")
            left.set_ylabel("Household load (kW)")
            left.set_xticks(np.arange(0, 24, 2))
            left.set_xlim(0, 23)

            right = left.twinx()
            right.plot(
                hours,
                arrays["price"],
                linestyle=":",
                linewidth=1.8,
                label="Wholesale price",
            )
            right.plot(
                hours,
                arrays["active_rate"][:, idx],
                marker="s",
                markersize=3,
                linewidth=1.5,
                label="Active nominal incentive rate",
            )
            right.set_ylabel("Rate (cent/kWh)")

            handles_left, labels_left = left.get_legend_handles_labels()
            handles_right, labels_right = right.get_legend_handles_labels()
            left.legend(
                handles_left + handles_right,
                labels_left + labels_right,
                fontsize=9,
                loc="upper left",
            )
            fig.tight_layout()
            fig.savefig(
                self.output_dir / f"household_{house_id}_profile_day_{day}.png",
                dpi=300,
                bbox_inches="tight",
            )
            plt.close(fig)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model_path", required=True)
    parser.add_argument("--config_path", required=True)
    parser.add_argument("--period", choices=["july", "august", "november"])
    parser.add_argument("--test_start", type=int)
    parser.add_argument("--test_end", type=int)
    parser.add_argument("--plot_days", nargs="*", type=int)
    parser.add_argument("--output_dir")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--no_plots", action="store_true")
    return parser.parse_args()


def resolve_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def main() -> None:
    args = parse_args()
    config_path = resolve_path(args.config_path)
    cfg = load_config(str(config_path))
    seed = int(cfg["general"]["seed"])

    if args.period:
        period_cfg = cfg["testing"]["periods"][args.period]
        start_day, end_day = map(int, period_cfg["range"])
        plot_days = [int(v) for v in period_cfg.get("plot_days", [])]
        period = args.period
    else:
        if args.test_start is None or args.test_end is None:
            raise ValueError("Supply --period or both --test_start and --test_end.")
        start_day, end_day = int(args.test_start), int(args.test_end)
        plot_days = []
        period = f"days_{start_day}_{end_day}"

    if args.plot_days is not None and len(args.plot_days) > 0:
        plot_days = [int(v) for v in args.plot_days]

    output_dir = (
        resolve_path(args.output_dir)
        if args.output_dir
        else PROJECT_ROOT / "results" / "benchmark" / f"seed_{seed}" / period
    )
    tester = BenchmarkTester(
        model_path=resolve_path(args.model_path),
        cfg=cfg,
        period=period,
        start_day=start_day,
        end_day=end_day,
        plot_days=plot_days,
        output_dir=output_dir,
        overwrite=args.overwrite,
        make_plots=not args.no_plots,
    )
    tester.run()


if __name__ == "__main__":
    main()
