"""Aggregate the three benchmark seeds into a paper-analysis workbook."""

from __future__ import annotations

import argparse
import itertools
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


EPS = 1.0e-12
DEFAULT_PERIODS = ("july", "august", "november")
DEFAULT_SEEDS = (0, 1, 2)


def safe_spearman(a: np.ndarray, b: np.ndarray) -> float:
    """Spearman correlation without SciPy (Pearson correlation of average ranks)."""
    left = pd.Series(np.asarray(a, dtype=float).ravel()).rank(method="average")
    right = pd.Series(np.asarray(b, dtype=float).ravel()).rank(method="average")
    value = left.corr(right, method="pearson")
    return float(value) if pd.notna(value) else float("nan")


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        sampled_max_row = min(sheet.max_row, 250)
        for column in sheet.iter_cols(
            min_row=1, max_row=sampled_max_row, min_col=1, max_col=sheet.max_column
        ):
            width = min(
                max(len(str(cell.value)) if cell.value is not None else 0 for cell in column) + 2,
                45,
            )
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(path)


def load_monthly(results_root: Path, period: str, seed: int) -> pd.DataFrame:
    path = results_root / f"seed_{seed}" / period / "results.xlsx"
    if not path.exists():
        raise FileNotFoundError(f"Missing benchmark result workbook: {path}")
    frame = pd.read_excel(path, sheet_name="Monthly")
    if len(frame) != 1:
        raise ValueError(f"Monthly sheet must contain exactly one row: {path}")
    return frame


def load_hourly(results_root: Path, period: str, seed: int) -> pd.DataFrame:
    path = results_root / f"seed_{seed}" / period / "results.xlsx"
    frame = pd.read_excel(path, sheet_name="Hourly_Aggregate")
    frame = frame.sort_values(["day", "hour"]).reset_index(drop=True)
    return frame


def signal_rows(period: str, hourly_by_seed: dict[int, pd.DataFrame]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for seed_a, seed_b in itertools.combinations(sorted(hourly_by_seed), 2):
        a = hourly_by_seed[seed_a]
        b = hourly_by_seed[seed_b]
        key_cols = ["day", "hour"]
        if not a[key_cols].equals(b[key_cols]):
            raise ValueError(f"Hourly alignment mismatch for {period}, seeds {seed_a}-{seed_b}.")

        raw_cols = [column for column in a.columns if column.startswith("raw_action_house_")]
        rate_cols = [
            column
            for column in a.columns
            if column.startswith("active_nominal_rate_house_")
        ]
        a_raw = a[raw_cols].to_numpy(dtype=float)
        b_raw = b[raw_cols].to_numpy(dtype=float)
        a_rate = a[rate_cols].to_numpy(dtype=float)
        b_rate = b[rate_cols].to_numpy(dtype=float)
        a_household_active = a_raw > EPS
        b_household_active = b_raw > EPS
        a_active = np.any(a_household_active, axis=1)
        b_active = np.any(b_household_active, axis=1)
        no_need = a["no_need_flag_external"].to_numpy(dtype=bool)
        overload = a["baseline_violation_flag_external"].to_numpy(dtype=bool)

        zero_nonzero_agreement = float(np.mean(a_active == b_active))
        household_zero_nonzero_agreement = float(
            np.mean(a_household_active == b_household_active)
        )
        exact_joint = float(
            np.mean(a["action_index"].to_numpy(dtype=int) == b["action_index"].to_numpy(dtype=int))
        )
        no_need_agreement = (
            float(np.mean((~a_active[no_need]) & (~b_active[no_need])))
            if np.any(no_need)
            else float("nan")
        )
        overload_active_agreement = (
            float(np.mean(a_active[overload] & b_active[overload]))
            if np.any(overload)
            else float("nan")
        )
        rows.append(
            {
                "period": period,
                "seed_pair": f"{seed_a}-{seed_b}",
                "zero_nonzero_agreement_percent": 100.0 * zero_nonzero_agreement,
                "household_zero_nonzero_agreement_percent": 100.0 * household_zero_nonzero_agreement,
                "exact_joint_action_agreement_percent": 100.0 * exact_joint,
                "raw_action_spearman": safe_spearman(a_raw, b_raw),
                "raw_action_MAE": float(np.mean(np.abs(a_raw - b_raw))),
                "active_nominal_rate_spearman": safe_spearman(a_rate, b_rate),
                "active_nominal_rate_MAE_cent_per_kWh": float(np.mean(np.abs(a_rate - b_rate))),
                "no_need_joint_zero_agreement_percent_external": 100.0 * no_need_agreement,
                "overload_both_active_agreement_percent_external": 100.0 * overload_active_agreement,
            }
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--results_root",
        default="results/benchmark",
        help="Benchmark results root relative to the project root or absolute.",
    )
    parser.add_argument(
        "--output",
        default="results/benchmark/summary.xlsx",
        help="Output workbook path relative to the project root or absolute.",
    )
    args = parser.parse_args()

    results_root = Path(args.results_root)
    if not results_root.is_absolute():
        results_root = PROJECT_ROOT / results_root
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = PROJECT_ROOT / output_path
    output_path.parent.mkdir(parents=True, exist_ok=True)

    seed_frames: list[pd.DataFrame] = []
    stability: list[dict[str, Any]] = []
    for period in DEFAULT_PERIODS:
        hourly_by_seed: dict[int, pd.DataFrame] = {}
        for seed in DEFAULT_SEEDS:
            monthly = load_monthly(results_root, period, seed).copy()
            monthly["period"] = period
            monthly["seed"] = seed
            seed_frames.append(monthly)
            hourly_by_seed[seed] = load_hourly(results_root, period, seed)
        stability.extend(signal_rows(period, hourly_by_seed))

    seed_performance = pd.concat(seed_frames, ignore_index=True)
    seed_performance = seed_performance.drop(columns=["day"], errors="ignore")

    numeric_columns = [
        column
        for column in seed_performance.select_dtypes(include=[np.number]).columns
        if column not in {"seed"}
    ]
    mean_sd_rows = []
    for period, frame in seed_performance.groupby("period", sort=False):
        for metric in numeric_columns:
            values = pd.to_numeric(frame[metric], errors="coerce").dropna()
            if values.empty:
                continue
            mean_sd_rows.append(
                {
                    "period": period,
                    "metric": metric,
                    "mean": float(values.mean()),
                    "sample_SD": float(values.std(ddof=1)) if len(values) > 1 else float("nan"),
                    "minimum": float(values.min()),
                    "maximum": float(values.max()),
                    "n_seeds": int(len(values)),
                }
            )
    mean_sd = pd.DataFrame(mean_sd_rows)
    stability_df = pd.DataFrame(stability)

    definitions = pd.DataFrame(
        [
            {
                "field": "Capacity role",
                "definition": (
                    "The 7 kW threshold is absent from benchmark training and policy selection; "
                    "capacity metrics are post-hoc external evaluations."
                ),
            },
            {
                "field": "Mean and SD",
                "definition": "Arithmetic mean and sample standard deviation across seeds 0, 1 and 2.",
            },
            {
                "field": "Signal stability",
                "definition": "Pairwise comparisons use identically aligned household-hour decisions.",
            },
            {
                "field": "Rebound",
                "definition": "N/A for benchmark because the response model is curtailment-only.",
            },
        ]
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        seed_performance.to_excel(writer, sheet_name="Seed_Performance", index=False)
        mean_sd.to_excel(writer, sheet_name="Mean_SD", index=False)
        stability_df.to_excel(writer, sheet_name="Signal_Stability", index=False)
        definitions.to_excel(writer, sheet_name="Definitions", index=False)
    style_workbook(output_path)
    print(f"Benchmark summary saved: {output_path}")


if __name__ == "__main__":
    main()
