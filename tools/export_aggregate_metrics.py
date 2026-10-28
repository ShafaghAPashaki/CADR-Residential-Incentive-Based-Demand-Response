"""Export seed-level monthly aggregate metrics behind the manuscript tables.

Reads the evaluation workbooks written by test.py (CADR, CBSP) and
benchmark/test_benchmark.py (EBDR) for July and August, and writes one long
CSV with one row per (method, seed, month, metric). Only monthly totals,
monthly means of daily statistics, and per-EU monthly totals are exported; no
hourly or daily household series are written.

The hourly workbooks contain licensed household data and are not distributed.
Re-create them by evaluating the released checkpoints with licensed inputs,
then run, from the project root:

    python tools/export_aggregate_metrics.py --out repro_runs/per_seed_monthly_metrics.csv

Use --cadr-root/--cbsp-root/--ebdr-root/--sensitivity-root to point at fresh
evaluation folders. Existing outputs are never overwritten.
"""
from __future__ import annotations

import argparse
import csv
import statistics as st
from pathlib import Path

import openpyxl

CAP_KW = 7.0
EXCEED_TOL_KW = 1e-6
MONTHS = ("july", "august")
SEEDS = (0, 1, 2)
HOUSES = ("661", "3039", "8565")


def _sheet(path: Path, name: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[name].iter_rows(values_only=True))
    finally:
        wb.close()
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if r[0] is not None]


def _daily(values_by_day: dict) -> tuple[float, float]:
    peaks, pars = [], []
    for day in sorted(values_by_day):
        v = values_by_day[day]
        if len(v) != 24:
            raise ValueError(f"day {day} has {len(v)} hours")
        mean = sum(v) / 24.0
        peaks.append(max(v))
        pars.append(max(v) / mean)
    return st.mean(peaks), st.mean(pars)


def cadr_like(path: Path, cap_kw: float = CAP_KW) -> dict:
    """Metrics from a test.py audit workbook (CADR, CBSP and the sensitivity runs)."""
    hourly = _sheet(path, "Hourly_Audit")
    daily = _sheet(path, "Daily_Summary")
    base, after = {}, {}
    for r in hourly:
        base.setdefault(r["day"], []).append(r["baseline_total"])
        after.setdefault(r["day"], []).append(r["after_total"])
    peak, par = _daily(after)
    base_peak, base_par = _daily(base)
    n_days = len(after)
    m = {
        "baseline_exceedance_hours": sum(r["baseline_total"] > cap_kw + EXCEED_TOL_KW for r in hourly),
        "baseline_cumulative_excess_kWh": sum(max(0.0, r["baseline_total"] - cap_kw) for r in hourly),
        "baseline_mean_load_kW": sum(r["baseline_total"] for r in hourly) / len(hourly),
        "baseline_mean_daily_peak_kW": base_peak,
        "baseline_daily_PAR": base_par,
        "exceedance_hours": sum(r["after_total"] > cap_kw + EXCEED_TOL_KW for r in hourly),
        "cumulative_excess_kWh": sum(max(0.0, r["after_total"] - cap_kw) for r in hourly),
        "mean_load_kW": sum(r["after_total"] for r in hourly) / len(hourly),
        "mean_daily_peak_kW": peak,
        "daily_PAR": par,
        "useful_DR_kWh": sum(r["useful_settlement_DR_energy"] for r in hourly),
        "settled_DR_kWh": sum(r["settlement_DR_energy"] for r in hourly),
        "net_curtailment_kWh": sum(r["true_hourly_curtailment"] for r in hourly),
        "shifted_out_kWh": sum(r["shifted_out_this_hour"] for r in hourly),
        "shifted_in_kWh": sum(r["shifted_in_this_hour"] for r in hourly),
        "unserved_EV_kWh": sum(d["unserved_energy_kwh"] for d in daily),
        "incentive_payment_cent": sum(r["incentive_payment"] for r in hourly),
        "SP_economic_value_cent": sum(r["SP_settlement_value"] for r in hourly),
        "SP_wholesale_value_cent": sum(r["SP_settlement_wholesale_value"] for r in hourly),
        "EU_utility_cent": sum(r["weighted_customer_utility"] for r in hourly),
        "incentive_offer_hours": sum(
            any((r[f"incentive_house_{h}"] or 0.0) > 1e-12 for h in HOUSES) for r in hourly
        ),
        "no_need_offer_hours": sum(bool(r["unnecessary_incentive"]) for r in hourly),
        "newly_created_violation_hours": sum(bool(r["newly_created_capacity_violation"]) for r in hourly),
        "rebound_induced_violation_hours": sum(bool(r["rebound_induced_violation"]) for r in hourly),
        "days": n_days,
    }
    total_settled = sum(r[f"settlement_DR_house_{h}"] for r in hourly for h in HOUSES)
    for i, h in enumerate(HOUSES, start=1):
        settled = sum(r[f"settlement_DR_house_{h}"] for r in hourly)
        payment = sum(r[f"incentive_payment_house_{h}"] for r in hourly)
        utility = sum(r[f"weighted_customer_utility_house_{h}"] for r in hourly)
        m[f"EU{i}_settled_DR_kWh"] = settled
        m[f"EU{i}_incentive_payment_cent"] = payment
        m[f"EU{i}_utility_cent"] = utility
        m[f"EU{i}_DR_share_percent"] = 100.0 * settled / total_settled if total_settled else 0.0
        m[f"EU{i}_participation_hours_percent"] = (
            100.0 * sum(r[f"settlement_DR_house_{h}"] > 1e-9 for r in hourly) / len(hourly)
        )
    # Appliance decomposition (Table 4) from the paper-ready workbook beside the audit.
    wb = openpyxl.load_workbook(path.with_name("results.xlsx"), read_only=True, data_only=True)
    for r in wb["Table_2_Appliances"].iter_rows(values_only=True):
        if r and r[0] in ("AC", "EV", "WM", "DW", "DR"):
            name = "Dryer" if r[0] == "DR" else r[0]
            m[f"{name}_curtailment_kWh"] = r[4]
            m[f"{name}_shifted_out_kWh"] = r[5]
            m[f"{name}_shifted_in_kWh"] = r[6]
    wb.close()
    return m


def ebdr(path: Path) -> dict:
    """Metrics from a benchmark/test_benchmark.py results workbook."""
    hourly = _sheet(path, "Hourly_Aggregate")
    after = {}
    for r in hourly:
        after.setdefault(r["day"], []).append(r["post_DR_load_kW"])
    peak, par = _daily(after)
    return {
        "exceedance_hours": sum(r["post_DR_load_kW"] > CAP_KW + EXCEED_TOL_KW for r in hourly),
        "cumulative_excess_kWh": sum(max(0.0, r["post_DR_load_kW"] - CAP_KW) for r in hourly),
        "mean_load_kW": sum(r["post_DR_load_kW"] for r in hourly) / len(hourly),
        "mean_daily_peak_kW": peak,
        "daily_PAR": par,
        "useful_DR_kWh": sum(r["useful_DR_kWh_external"] for r in hourly),
        "net_curtailment_kWh": sum(r["total_curtailment_kWh"] for r in hourly),
        "incentive_payment_cent": sum(r["incentive_payment_cent"] for r in hourly),
        "incentive_offer_hours": sum(bool(r["intervention_flag"]) for r in hourly),
        "no_need_offer_hours": sum(
            bool(r["intervention_flag"]) and bool(r["no_need_flag_external"]) for r in hourly
        ),
        "days": len(after),
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--cadr-root", default="results")
    ap.add_argument("--cbsp-root", default="results/matched_hems_ablation")
    ap.add_argument("--ebdr-root", default="results/benchmark")
    ap.add_argument("--sensitivity-root", default="results/capacity_incentive_sensitivity",
                    help="Folder holding cap6/, cap8/ and lev10/ July evaluations (Appendix H/I).")
    ap.add_argument("--skip-sensitivity", action="store_true")
    ap.add_argument("--out", default="repro_runs/per_seed_monthly_metrics.csv")
    args = ap.parse_args()
    out = Path(args.out)
    if out.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {out}")

    rows = []
    for seed in SEEDS:
        for month in MONTHS:
            sources = (
                ("CADR", Path(args.cadr_root) / f"seed_{seed}" / month / "audit.xlsx", cadr_like),
                ("CBSP", Path(args.cbsp_root) / f"seed_{seed}" / month / "audit.xlsx", cadr_like),
                ("EBDR", Path(args.ebdr_root) / f"seed_{seed}" / month / "results.xlsx", ebdr),
            )
            for method, path, fn in sources:
                if not path.exists():
                    raise FileNotFoundError(f"{method} workbook not found: {path}")
                for metric, value in fn(path).items():
                    rows.append({"method": method, "seed": seed, "month": month,
                                 "metric": metric, "value": value})

    if not args.skip_sensitivity:
        # Appendix H (capacity target) and Appendix I (ten incentive levels), Seed 0, July.
        # The 7-kW / five-level reference is the CADR Seed-0 July row above.
        for run, cap in (("cap6", 6.0), ("cap8", 8.0), ("lev10", 7.0)):
            path = Path(args.sensitivity_root) / run / "july" / "audit.xlsx"
            if not path.exists():
                raise FileNotFoundError(f"{run} workbook not found: {path}")
            for metric, value in cadr_like(path, cap).items():
                rows.append({"method": f"CADR_{run}", "seed": 0, "month": "july",
                             "metric": metric, "value": value})

    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("x", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["method", "seed", "month", "metric", "value"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {len(rows)} rows to {out}")


if __name__ == "__main__":
    main()
